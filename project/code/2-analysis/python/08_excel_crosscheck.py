"""Langkah 10 (bagian Excel): hitung ulang angka kunci dengan rumus Excel (COUNTIFS), dihitung ulang LibreOffice, bandingkan dengan SQL/Python.
File xlsx memuat salinan data sumber, jadi tidak di-commit (dibuat ulang skrip ini)."""
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "1-data-preparation" / "python"))
from importlib import import_module  # noqa: E402

from common import RES_ANALYSIS  # noqa: E402

df, _, _ = import_module("01_load_audit").load_raw()
fr = pd.read_csv(RES_ANALYSIS / "final_rule.csv").iloc[0]
th = {k: float(v) for k, v in re.findall(r"([LMH])=(\d+)", fr.threshold)}
dd, rr = (float(x) for x in re.search(r"suhu-rpm: ([\d.]+) K, (\d+) rpm", fr.threshold).groups())
tq = float(re.search(r"torque: (\d+) Nm", fr.threshold).group(1))
wb = Workbook()
ws = wb.active; ws.title = "Data"
ws.append(["udi", "type", "air", "proc", "rpm", "torque", "wear", "failure", "twf", "hdf", "pwf", "osf", "rnf", "temp_diff", "wear_x_torque", "alert_final", "alert_wear200"])
for i, r in enumerate(df.itertuples(), start=2):
    ws.append([r.udi, r.type, r.air_temp_k, r.process_temp_k, r.rpm, r.torque_nm, r.tool_wear_min, r.machine_failure, r.twf, r.hdf, r.pwf, r.osf, r.rnf,
               f"=D{i}-C{i}", f"=G{i}*F{i}",
               f'=IF(OR(O{i}>=IF(B{i}="L",Check!$B$2,IF(B{i}="M",Check!$B$3,Check!$B$4)),AND(N{i}<Check!$B$5,E{i}<Check!$B$6),F{i}>=Check!$B$7),1,0)', f"=IF(G{i}>=200,1,0)"])
ck = wb.create_sheet("Check")
for i, (k, v) in enumerate([("ambang wear x torque L", th["L"]), ("ambang wear x torque M", th["M"]), ("ambang wear x torque H", th["H"]), ("selisih suhu <", dd), ("rpm <", rr), ("torque >=", tq)], start=2):
    ck.cell(i, 1, k); ck.cell(i, 2, v)
R = 10  # baris mulai pemeriksaan
checks = []
def add(name, formula):
    global R
    ck.cell(R, 1, name); ck.cell(R, 2, formula); checks.append((name, R)); R += 1
N = len(df) + 1
for t in "LMH":
    add(f"run tipe {t}", f'=COUNTIFS(Data!B2:B{N},"{t}")'); add(f"gagal tipe {t}", f'=COUNTIFS(Data!B2:B{N},"{t}",Data!H2:H{N},1)')
add("gagal semua", f"=SUM(Data!H2:H{N})")
for k, col in (("TP", "1,Data!H2:H%d,1" % N), ("FP", "1,Data!H2:H%d,0" % N), ("FN", "0,Data!H2:H%d,1" % N), ("TN", "0,Data!H2:H%d,0" % N)):
    add(f"tool wear >= 200: {k}", f"=COUNTIFS(Data!Q2:Q{N},{col})")
for k, col in (("TP", "1,Data!H2:H%d,1" % N), ("FP", "1,Data!H2:H%d,0" % N), ("FN", "0,Data!H2:H%d,1" % N), ("TN", "0,Data!H2:H%d,0" % N)):
    add(f"aturan terpilih (semua data): {k}", f"=COUNTIFS(Data!P2:P{N},{col})")
tmp = Path(tempfile.mkdtemp())
src = tmp / "crosscheck_src.xlsx"
wb.save(src)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(tmp / "out"), str(src)], check=True, capture_output=True, timeout=280)
v = load_workbook(tmp / "out" / "crosscheck_src.xlsx", data_only=True)["Check"]
pd_t = {k: (int(((df.type == k)).sum()), int(((df.type == k) & (df.machine_failure == 1)).sum())) for k in "LMH"}
mask_w = (df.tool_wear_min >= 200).values
y = df.machine_failure.values
fm = np.where(df.type == "L", th["L"], np.where(df.type == "M", th["M"], th["H"]))
mask_f = ((df.tool_wear_min * df.torque_nm).values >= fm) | (((df.process_temp_k - df.air_temp_k).values < dd) & (df.rpm.values < rr)) | (df.torque_nm.values >= tq)
conf = lambda m: [int((m & (y == 1)).sum()), int((m & (y == 0)).sum()), int((~m & (y == 1)).sum()), int((~m & (y == 0)).sum())]
exp = {}
for t in "LMH":
    exp[f"run tipe {t}"], exp[f"gagal tipe {t}"] = pd_t[t]
exp["gagal semua"] = int(y.sum())
for k, val in zip(("TP", "FP", "FN", "TN"), conf(mask_w)):
    exp[f"tool wear >= 200: {k}"] = val
for k, val in zip(("TP", "FP", "FN", "TN"), conf(mask_f)):
    exp[f"aturan terpilih (semua data): {k}"] = val
out = []
for name, row in checks:
    xv = v.cell(row, 2).value
    out.append((name, exp[name], xv, "PASS" if xv is not None and int(round(xv)) == exp[name] else "FAIL"))
pd.DataFrame(out, columns=["pemeriksaan", "nilai_sql_python", "nilai_excel", "status"]).to_csv(RES_ANALYSIS / "crosscheck_excel.csv", index=False)
# hasil rekalkulasi LibreOffice disimpan sebagai crosscheck.xlsx (tidak di-commit)
shutil.copy(tmp / "out" / "crosscheck_src.xlsx", RES_ANALYSIS / "crosscheck.xlsx")
shutil.rmtree(tmp)
assert all(o[3] == "PASS" for o in out), [o for o in out if o[3] != "PASS"]
fra = pd.read_csv(RES_ANALYSIS / "final_rule.csv").iloc[0]
assert [fra.tp_all, fra.fp_all, fra.fn_all, fra.tn_all] == conf(mask_f), "aturan terpilih berbeda dari final_rule.csv"
print(len(out), "pemeriksaan Excel PASS")
